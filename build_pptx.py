#!/usr/bin/env python3
"""build_pptx.py — Étape 2b : storyboard.json -> presentation.pptx.

Lit `llm-output/storyboard.json` (produit par `llm-process.py`) et génère un
PowerPoint 16:9. Chaque slide reçoit son titre, ses puces, son éventuelle image
(chemin résolu depuis la racine du projet) et son voiceover en notes du
présentateur.

Avec `--audio`, un fichier audio TTS est généré depuis le voiceover de chaque
slide et intégré au PPT en **lecture automatique** : en diaporama, avancer ou
reculer déclenche le voiceover de la slide affichée.

Deux moteurs TTS (`--tts`) :

- `say` (défaut) : commande macOS, gratuit et hors-ligne, puis `ffmpeg` pour
  convertir en mp3. Voix via `--voice` (ex. `Thomas`, `Amélie`).
- `elevenlabs` : voix humaines très naturelles. Nécessite une clé API dans
  `ELEVENLABS_API_KEY` (env ou `config.json`). Voix via `--voice` (nom ou id) ;
  `--list-voices` affiche les voix disponibles du compte.

Usage :
    python build_pptx.py [--in llm-output] [--out llm-output/presentation.pptx]
                         [--audio] [--tts say|elevenlabs] [--voice Thomas]
                         [--speed 1.25] [--rate 180]
                         [--audio-dir llm-output/audio]
                         [--keep-audio] [--config config.json]
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x44, 0x44, 0x44)

ELEVENLABS_API = "https://api.elevenlabs.io/v1"
ELEVENLABS_MODEL = "eleven_multilingual_v2"
ELEVENLABS_VOICE = "George"  # voix par défaut si --voice n'est pas fourni
ELEVENLABS_SPEED = 1.25      # vitesse par défaut pour ElevenLabs
SAY_VOICE = "Thomas"
PREVIEW_TEXT = (
    "Bonjour, je suis votre présentateur. Voici un exemple de voix pour "
    "expliquer une notion scientifique de façon claire et naturelle."
)


def load_storyboard(path: Path) -> dict:
    text = path.read_text(encoding="utf-8").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.MULTILINE).strip()
    return json.loads(text)


# ---------------------------------------------------------------------------
# Voiceover TTS + lecture automatique
# ---------------------------------------------------------------------------

def synthesize_say(text: str, dest: Path, voice: str, rate: int) -> Path | None:
    """Génère `dest` (mp3) depuis `text` via `say` + `ffmpeg`."""
    if not text.strip():
        return None
    if not shutil.which("say"):
        raise SystemExit("TTS indisponible : la commande macOS 'say' est requise.")
    if not shutil.which("ffmpeg"):
        raise SystemExit("ffmpeg est requis pour convertir l'audio en mp3.")
    dest.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        aiff = Path(tmp) / "vo.aiff"
        subprocess.run(
            ["say", "-v", voice, "-r", str(rate), "-o", str(aiff), text], check=True
        )
        subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error", "-i", str(aiff),
             "-ac", "1", "-b:a", "64k", str(dest)],
            check=True,
        )
    return dest


def _elevenlabs_http(url: str, api_key: str, payload: dict | None = None) -> bytes:
    """Appel HTTP brut à l'API ElevenLabs (stdlib, sans dépendance)."""
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(
        url, data=data, method="POST" if data is not None else "GET"
    )
    req.add_header("xi-api-key", api_key)
    req.add_header("accept", "audio/mpeg")
    if data is not None:
        req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            return resp.read()
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")
        raise SystemExit(f"ElevenLabs HTTP {exc.code} : {detail}")
    except urllib.error.URLError as exc:
        raise SystemExit(f"ElevenLabs injoignable : {exc.reason}")


def list_elevenlabs_voices(api_key: str) -> list[dict]:
    raw = _elevenlabs_http(f"{ELEVENLABS_API}/voices", api_key)
    return json.loads(raw).get("voices", [])


def voice_labels(v: dict) -> tuple[str, str]:
    """Renvoie (nom complet, nom court) ; gère « George - Storyteller »."""
    name = v.get("name", "").lower()
    return name, name.split(" - ")[0].strip()


def find_voice(voices: list[dict], query: str) -> dict | None:
    """Trouve une voix par id, nom exact, nom court (avant ' - ') ou préfixe."""
    q = query.strip().lower()
    for v in voices:
        if v.get("voice_id") == query:
            return v
    for v in voices:
        if q in voice_labels(v):
            return v
    for v in voices:
        name, short = voice_labels(v)
        if short.startswith(q) or q in name:
            return v
    return None


def resolve_elevenlabs_voice(api_key: str, voice: str) -> str:
    """Accepte un voice_id (20 car.) ou un nom de voix, résolu via l'API."""
    if re.fullmatch(r"[A-Za-z0-9]{20}", voice):
        return voice
    voices = list_elevenlabs_voices(api_key)
    v = find_voice(voices, voice)
    if v is not None:
        return v["voice_id"]
    names = ", ".join(sorted({voice_labels(v)[1] for v in voices}))
    raise SystemExit(
        f"Voix ElevenLabs « {voice} » introuvable.\nDisponibles : {names}\n"
        "Astuce : python build_pptx.py --tts elevenlabs --list-voices"
    )


def synthesize_elevenlabs(
    text: str, dest: Path, voice: str, model_id: str, api_key: str
) -> Path | None:
    """Génère `dest` (mp3) depuis `text` via l'API ElevenLabs."""
    if not text.strip():
        return None
    if not api_key:
        raise SystemExit(
            "Clé ElevenLabs manquante : renseigne ELEVENLABS_API_KEY dans\n"
            "l'environnement, ou dans config.json (champ env), puis relance."
        )
    voice_id = resolve_elevenlabs_voice(api_key, voice)
    url = f"{ELEVENLABS_API}/text-to-speech/{voice_id}?output_format=mp3_44100_128"
    payload = {
        "text": text,
        "model_id": model_id,
        "voice_settings": {
            "stability": 0.5,
            "similarity_boost": 0.75,
            "style": 0.0,
            "use_speaker_boost": True,
        },
    }
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(_elevenlabs_http(url, api_key, payload))
    return dest


def synthesize(text: str, dest: Path, cfg: dict) -> Path | None:
    """Dispatch TTS selon `cfg['tts']`."""
    if not text.strip():
        return None
    if cfg["tts"] == "elevenlabs":
        return synthesize_elevenlabs(
            text, dest, cfg["voice"], cfg["model_id"], cfg.get("api_key", "")
        )
    return synthesize_say(text, dest, cfg["voice"], cfg["rate"])


def play_audio(path: Path) -> None:
    """Joue un audio localement si un lecteur est disponible (afplay/ffplay)."""
    if shutil.which("afplay"):
        subprocess.run(["afplay", str(path)])
    elif shutil.which("ffplay"):
        subprocess.run(
            ["ffplay", "-nodisp", "-autoexit", "-loglevel", "quiet", str(path)]
        )


def atempo_filter(speed: float) -> str:
    """Filtres ffmpeg atempo pour une vitesse arbitraire (hauteur conservée)."""
    parts = []
    s = speed
    while s > 2.0:
        parts.append("atempo=2.0")
        s /= 2.0
    while s < 0.5:
        parts.append("atempo=0.5")
        s /= 0.5
    parts.append(f"atempo={s:.4f}")
    return ",".join(parts)


def apply_speed(path: Path, speed: float) -> None:
    """Accélère/ralentit un mp3 sans changer la voix (ffmpeg atempo)."""
    if abs(speed - 1.0) < 1e-3:
        return
    if not shutil.which("ffmpeg"):
        print("  ffmpeg absent : vitesse non appliquée.")
        return
    tmp = path.with_name(path.stem + ".speed" + path.suffix)
    subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error", "-i", str(path),
         "-filter:a", atempo_filter(speed), "-b:a", "128k", str(tmp)],
        check=True,
    )
    tmp.replace(path)


def set_autoplay(shape) -> None:
    """Passe le média de la forme en lecture automatique (delay=0).

    Voir https://github.com/scanny/python-pptx/issues/427 — python-pptx
    n'expose pas ce réglage, on édite le nœud de timing.
    """
    pic = shape._element
    cNvPr = pic.find(".//" + qn("p:cNvPr"))
    if cNvPr is None:
        return
    spid = cNvPr.get("id")
    root = pic.getroottree().getroot()
    for tgt in root.iter(qn("p:spTgt")):
        if tgt.get("spid") != spid:
            continue
        cmedia = tgt.getparent().getparent()
        cond = cmedia.find(".//" + qn("p:cond"))
        if cond is not None:
            cond.set("delay", "0")
        return


def embed_audio(slide, mp3: Path) -> None:
    shape = slide.shapes.add_movie(
        str(mp3), Inches(-2), Inches(0), Inches(1), Inches(1), mime_type="audio/mpeg"
    )
    set_autoplay(shape)


# ---------------------------------------------------------------------------
# Rendu des slides
# ---------------------------------------------------------------------------

def add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), SLIDE_W - Inches(1.2), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.18), SLIDE_W - Inches(1.2), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_bullets(slide, bullets: list[str], left, top, width, height, size: int) -> None:
    if not bullets:
        return
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {b}"
        run.font.size = Pt(size)
        run.font.color.rgb = GREY


def add_image(slide, img_path: Path, left, top, max_w, max_h) -> bool:
    if not img_path.is_file():
        print(f"  image introuvable, ignorée : {img_path}")
        return False
    with Image.open(img_path) as im:
        ratio = im.width / im.height
    w, h = max_w, int(max_w / ratio)
    if h > max_h:
        h, w = max_h, int(max_h * ratio)
    slide.shapes.add_picture(
        str(img_path), int(left + (max_w - w) / 2), int(top + (max_h - h) / 2), width=w, height=h
    )
    return True


def add_caption(slide, text: str, left, top, width) -> None:
    box = slide.shapes.add_textbox(left, top, width, Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = GREY


def add_notes(slide, voiceover: str) -> None:
    if voiceover:
        slide.notes_slide.notes_text_frame.text = voiceover


def build(data: dict, out_path: Path, root: Path, audio_cfg: dict | None = None) -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    slides = data.get("slides", [])
    for i, s in enumerate(slides, start=1):
        layout = s.get("layout", "content")
        slide = prs.slides.add_slide(blank)

        if layout == "title":
            box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), SLIDE_W - Inches(1.6), Inches(1.6))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = s.get("title", "")
            run.font.size = Pt(40)
            run.font.bold = True
            run.font.color.rgb = ACCENT
            sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.1), SLIDE_W - Inches(1.6), Inches(1.4))
            stf = sub.text_frame
            stf.word_wrap = True
            for j, line in enumerate([data.get("source"), data.get("audience")]):
                if not line:
                    continue
                p = stf.paragraphs[0] if j == 0 else stf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = line
                r.font.size = Pt(18)
                r.font.color.rgb = GREY
        else:
            add_title(slide, s.get("title", ""))
            bullets = s.get("bullets") or []
            image = s.get("image")
            if image:
                add_bullets(
                    slide, bullets, Inches(0.6), Inches(1.6), Inches(5.6), Inches(5.2),
                    size=16 if len(bullets) > 5 else 18,
                )
                add_image(slide, root / image, Inches(6.5), Inches(1.5), Inches(6.3), Inches(5.0))
                if s.get("image_caption"):
                    add_caption(slide, s["image_caption"], Inches(6.5), Inches(6.6), Inches(6.3))
            else:
                add_bullets(
                    slide, bullets, Inches(0.8), Inches(1.6), SLIDE_W - Inches(1.6), Inches(5.2),
                    size=18 if len(bullets) <= 6 else 15,
                )

        voiceover = s.get("voiceover", "")
        add_notes(slide, voiceover)

        if audio_cfg and voiceover:
            mp3 = audio_cfg["dir"] / f"slide_{i:02d}.mp3"
            if audio_cfg.get("keep") and mp3.is_file():
                pass
            else:
                print(f"  TTS slide {i} ({audio_cfg['tts']})…")
                if synthesize(voiceover, mp3, audio_cfg):
                    apply_speed(mp3, audio_cfg["speed"])
            if mp3.is_file():
                embed_audio(slide, mp3)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return len(slides)


def load_env(path: Path) -> dict[str, str]:
    """Lit le champ `env` de config.json (clés API), sans échouer si absent."""
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    return {k: str(v) for k, v in (data.get("env") or {}).items() if v}


def main() -> None:
    ap = argparse.ArgumentParser(description="storyboard.json -> presentation.pptx")
    ap.add_argument("--in", dest="indir", type=Path, default=Path("llm-output"))
    ap.add_argument("--out", dest="outfile", type=Path, default=None)
    ap.add_argument("--audio", action="store_true",
                    help="génère un voiceover TTS par slide en lecture automatique")
    ap.add_argument("--tts", choices=["say", "elevenlabs"], default="say",
                    help="moteur TTS (défaut : say macOS)")
    ap.add_argument("--voice", default=None,
                    help="voix : nom/id ElevenLabs, ou voix macOS pour --tts say "
                         "(liste séparée par des virgules avec --preview)")
    ap.add_argument("--tts-model", default=ELEVENLABS_MODEL,
                    help="modèle ElevenLabs (ex. eleven_multilingual_v2, eleven_flash_v2_5)")
    ap.add_argument("--rate", type=int, default=180, help="débit say (mots/minute)")
    ap.add_argument("--speed", type=float, default=None,
                    help="vitesse de lecture (1.0 = normal ; défaut ElevenLabs : 1.25)")
    ap.add_argument("--audio-dir", type=Path, default=None,
                    help="dossier des mp3 (défaut : <in>/audio)")
    ap.add_argument("--keep-audio", action="store_true",
                    help="réutilise les mp3 existants au lieu de les régénérer")
    ap.add_argument("--config", type=Path, default=Path("config.json"),
                    help="config locale (champ env : clés API) — ignorée par Git")
    ap.add_argument("--list-voices", action="store_true",
                    help="avec --tts elevenlabs : liste les voix du compte et quitte")
    ap.add_argument("--preview", nargs="?", const=PREVIEW_TEXT, default=None,
                    metavar="TEXTE",
                    help="avec --tts elevenlabs : échantillon TTS de chaque voix "
                         "(défaut : phrase FR) puis lecture locale")
    args = ap.parse_args()

    root = Path.cwd()
    env_extra = load_env(args.config if args.config.is_absolute() else root / args.config)
    api_key = os.environ.get("ELEVENLABS_API_KEY") or env_extra.get("ELEVENLABS_API_KEY", "")
    indir = args.indir if args.indir.is_absolute() else root / args.indir

    if args.list_voices:
        if args.tts != "elevenlabs":
            raise SystemExit("--list-voices nécessite --tts elevenlabs")
        if not api_key:
            raise SystemExit("Clé ElevenLabs manquante (ELEVENLABS_API_KEY).")
        for v in list_elevenlabs_voices(api_key):
            full = v.get("name", "?")
            short = full.split(" - ")[0].strip()
            print(f"{short:<14} {v.get('voice_id', '?')}")
            if " - " in full:
                print(f"{'':<14} {full.split(' - ', 1)[1]}")
            url = v.get("preview_url")
            if url:
                print(f"{'':<14} écoute : {url}")
        return

    if args.preview is not None:
        if args.tts != "elevenlabs":
            raise SystemExit("--preview nécessite --tts elevenlabs")
        if not api_key:
            raise SystemExit("Clé ElevenLabs manquante (ELEVENLABS_API_KEY).")
        voices = list_elevenlabs_voices(api_key)
        wanted = [w.strip() for w in (args.voice or "").split(",") if w.strip()]
        selected = voices if not wanted else []
        for w in wanted:
            v = find_voice(voices, w)
            if v is None:
                raise SystemExit(f"Voix « {w} » introuvable.")
            selected.append(v)
        preview_dir = indir / "voice-preview"
        preview_dir.mkdir(parents=True, exist_ok=True)
        speed = args.speed if args.speed is not None else ELEVENLABS_SPEED
        print(f"{len(selected)} voix, texte : « {args.preview} » (vitesse {speed}x)\n")
        for v in selected:
            short = v.get("name", v.get("voice_id", "voix")).split(" - ")[0].strip()
            name = re.sub(r"[^A-Za-z0-9_-]+", "_", short)
            dest = preview_dir / f"{name}.mp3"
            synthesize_elevenlabs(args.preview, dest, v["voice_id"], args.tts_model, api_key)
            apply_speed(dest, speed)
            print(f"  {short:<14} -> {dest}")
            play_audio(dest)
        print(f"\néchantillons : {preview_dir}")
        return

    storyboard = indir / "storyboard.json"
    if not storyboard.is_file():
        raise SystemExit(f"Introuvable : {storyboard}")

    data = load_storyboard(storyboard)
    outfile = args.outfile or (indir / "presentation.pptx")
    if not outfile.is_absolute():
        outfile = root / outfile

    audio_cfg = None
    if args.audio:
        audio_dir = args.audio_dir or (indir / "audio")
        if not audio_dir.is_absolute():
            audio_dir = root / audio_dir
        voice = args.voice or (ELEVENLABS_VOICE if args.tts == "elevenlabs" else SAY_VOICE)
        speed = args.speed if args.speed is not None else (
            ELEVENLABS_SPEED if args.tts == "elevenlabs" else 1.0
        )
        audio_cfg = {
            "tts": args.tts,
            "dir": audio_dir,
            "voice": voice,
            "rate": args.rate,
            "speed": speed,
            "keep": args.keep_audio,
            "model_id": args.tts_model,
            "api_key": api_key,
        }

    n = build(data, outfile, root, audio_cfg)
    print(f"{n} slides -> {outfile}")
    if audio_cfg:
        print(f"voiceover -> {audio_cfg['dir']} "
              f"({audio_cfg['tts']}, voix {audio_cfg['voice']}, lecture automatique)")


if __name__ == "__main__":
    main()
